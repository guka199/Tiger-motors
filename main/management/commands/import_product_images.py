import os
import zipfile

from django.core.management.base import BaseCommand, CommandError
from django.core.files.base import ContentFile

from main.models import Product, ProductImage


class Command(BaseCommand):
    help = 'Import product images from an Excel (.xlsx) file'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', help='Path to the Excel file')
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Preview what would be imported without saving anything',
        )
        parser.add_argument(
            '--row-start',
            type=int,
            default=1,
            help='First data row number (default: 1). Use 2 if your sheet has a header row.',
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError('openpyxl is required. Install it with: pip install openpyxl')

        filepath = options['excel_file']
        dry_run = options['dry_run']
        row_start = options['row_start']

        if not os.path.exists(filepath):
            raise CommandError(f'File not found: {filepath}')

        if dry_run:
            self.stdout.write(self.style.WARNING('DRY RUN — nothing will be saved\n'))

        wb = load_workbook(filepath)
        ws = wb.active

        # Build: anchor_row (0-indexed) -> [(col_idx, img_ref), ...]
        # Images in Excel are "floating" drawings anchored to a cell.
        # anchor._from.row is 0-indexed (row 1 in Excel = 0 here).
        row_image_refs = {}
        for img in ws._images:
            anchor = img.anchor
            if not hasattr(anchor, '_from'):
                self.stdout.write(
                    self.style.WARNING(f'Skipping image with unsupported anchor type: {type(anchor).__name__}')
                )
                continue
            row_idx = anchor._from.row  # 0-indexed
            col_idx = anchor._from.col  # 0-indexed  (A=0, B=1, C=2)
            if col_idx > 2:             # only care about first 3 columns
                continue
            row_image_refs.setdefault(row_idx, []).append((col_idx, img.ref))

        # Sort images within each row by column so is_main goes to the leftmost one
        for row_idx in row_image_refs:
            row_image_refs[row_idx].sort(key=lambda x: x[0])

        # Pre-read all image binary data directly from the xlsx ZIP archive
        image_data = {}  # (row_idx, col_idx) -> bytes
        with zipfile.ZipFile(filepath, 'r') as zf:
            available = set(zf.namelist())
            for row_idx, img_list in row_image_refs.items():
                for col_idx, ref in img_list:
                    zip_path = self._resolve_zip_path(ref)
                    if zip_path in available:
                        image_data[(row_idx, col_idx)] = zf.read(zip_path)
                    else:
                        # Fallback: search by basename
                        basename = os.path.basename(ref)
                        matches = [f for f in available if f.endswith(basename)]
                        if matches:
                            image_data[(row_idx, col_idx)] = zf.read(matches[0])
                        else:
                            self.stdout.write(
                                self.style.WARNING(
                                    f'  Image not found in ZIP: {ref!r} (tried {zip_path!r})'
                                )
                            )

        total_products = 0
        total_images = 0
        total_skipped = 0

        for row_cells in ws.iter_rows(min_row=row_start):
            excel_row = row_cells[0].row
            anchor_row = excel_row - 1  # convert to 0-indexed to match anchor

            product_id_val = row_cells[0].value
            if not product_id_val:
                continue
            product_id = str(product_id_val).strip()

            imgs_for_row = [
                (col, image_data[(anchor_row, col)])
                for col, _ in row_image_refs.get(anchor_row, [])
                if (anchor_row, col) in image_data
            ]

            self.stdout.write(
                f'Row {excel_row}: product_id={product_id!r}  images found={len(imgs_for_row)}'
            )

            if dry_run:
                for col, data in imgs_for_row:
                    ext = self._detect_ext(data)
                    self.stdout.write(f'  -> col {col + 1}: {len(data)} bytes  (.{ext})')
                total_products += 1
                continue

            try:
                product = Product.objects.get(product_id=product_id)
            except Product.DoesNotExist:
                self.stdout.write(self.style.WARNING(f'  Product {product_id!r} not found — skipping'))
                total_skipped += 1
                continue

            for i, (col, data) in enumerate(imgs_for_row):
                ext = self._detect_ext(data)
                filename = f'{product_id}_col{col + 1}.{ext}'
                pi = ProductImage(product=product, is_main=(i == 0))
                pi.image.save(filename, ContentFile(data), save=True)
                self.stdout.write(self.style.SUCCESS(f'  Saved: {filename}'))
                total_images += 1

            total_products += 1

        self.stdout.write('\n' + '=' * 50)
        if dry_run:
            self.stdout.write(f'DRY RUN complete: {total_products} product rows found')
        else:
            self.stdout.write(
                self.style.SUCCESS(
                    f'Done: {total_products} products processed, '
                    f'{total_images} images saved, '
                    f'{total_skipped} products skipped (not found in DB)'
                )
            )

    def _resolve_zip_path(self, ref):
        """Convert an image relationship target to its path inside the xlsx ZIP."""
        if ref.startswith('../'):
            return 'xl' + ref[2:]          # '../media/image1.png' -> 'xl/media/image1.png'
        if ref.startswith('/'):
            return ref.lstrip('/')
        if not ref.startswith('xl/'):
            return f'xl/media/{os.path.basename(ref)}'
        return ref

    def _detect_ext(self, data):
        if data[:8] == b'\x89PNG\r\n\x1a\n':
            return 'png'
        if data[:3] == b'\xff\xd8\xff':
            return 'jpg'
        if data[:6] in (b'GIF87a', b'GIF89a'):
            return 'gif'
        if data[:4] == b'RIFF' and data[8:12] == b'WEBP':
            return 'webp'
        return 'jpg'
